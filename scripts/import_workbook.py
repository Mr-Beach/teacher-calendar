"""One-time conversion: pacing-guide .xlsx workbook -> plain course data file (JSON).

Reads the workbook's Days and Lessons sheets (the shape described in SPEC.md) and
writes a single JSON file that becomes the source of truth going forward. This is
a one-way migration, not a sync step: after running it, edit the JSON directly and
leave the workbook as an archive.

Usage:
    python3 scripts/import_workbook.py <path-to-workbook.xlsx> <course-slug> [--title "Math 6"] [--year "2026-27"]

Writes courses/<course-slug>.json
"""
import argparse
import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from engine import run_all_checks

DAYS_SHEET = "Days"
LESSONS_SHEET = "Lessons"

# Column indices (0-based) in the Days sheet, header row 2.
DAY_DATE, DAY_WEEKDAY, DAY_TYPE, DAY_NOTE = 0, 1, 2, 9

# Column indices (0-based) in the Lessons sheet, header row 2.
LSN_TOPIC, LSN_CODE, LSN_TITLE, LSN_KIND, LSN_STUDENT_TEXT, LSN_HOMEWORK = 1, 2, 3, 4, 5, 6


def read_school_days(ws):
    days = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        date = row[DAY_DATE]
        if date is None:
            continue
        days.append(
            {
                "date": date.date().isoformat(),
                "weekday": row[DAY_WEEKDAY],
                "type": row[DAY_TYPE],
                "note": row[DAY_NOTE],
            }
        )
    days.sort(key=lambda d: d["date"])
    return days


def read_sequence(ws):
    sequence = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        topic, code, title, kind = row[LSN_TOPIC], row[LSN_CODE], row[LSN_TITLE], row[LSN_KIND]
        if title is None and kind is None:
            continue
        sequence.append(
            {
                "topic": topic,
                "lesson_code": code,
                "district_title": title,
                "kind": kind,
                "student_text": row[LSN_STUDENT_TEXT],
                "homework": row[LSN_HOMEWORK],
            }
        )
    return sequence


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("slug", help="course slug, e.g. math6 -> courses/math6.json")
    parser.add_argument("--title", default=None, help="display name, e.g. 'Math 6'")
    parser.add_argument("--year", default=None, help="school year label, e.g. '2026-27'")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    for sheet in (DAYS_SHEET, LESSONS_SHEET):
        if sheet not in wb.sheetnames:
            sys.exit(f"expected a '{sheet}' sheet, found: {wb.sheetnames}")

    school_days = read_school_days(wb[DAYS_SHEET])
    sequence = read_sequence(wb[LESSONS_SHEET])

    instructional_days = sum(1 for d in school_days if d["type"] == "Instruction")
    print(f"school days: {len(school_days)} ({instructional_days} Instruction)", file=sys.stderr)
    print(f"sequence: {len(sequence)} lessons", file=sys.stderr)
    if instructional_days != len(sequence):
        print(
            f"NOTE: instructional day count ({instructional_days}) and sequence length "
            f"({len(sequence)}) differ — that's expected once you start editing day types.",
            file=sys.stderr,
        )

    course = {
        "course": args.title or args.slug,
        "school_year": args.year,
        "school_days": school_days,
        "sequence": sequence,
    }

    # Run the full checklist (PLANNING.md's "Sanity checks") right at import,
    # not just the leftover/instructional-day-count arithmetic above -- so a
    # fresh course starts from a known-checked state instead of surfacing
    # problems months into the school year.
    any_warnings = False
    for label, warnings in run_all_checks(course):
        if not warnings:
            continue
        any_warnings = True
        print(f"NOTE ({label}):", file=sys.stderr)
        for w in warnings:
            print(f"  {w}", file=sys.stderr)
    if not any_warnings:
        print("checks: none of the automated checks found anything to flag", file=sys.stderr)

    out_path = Path("courses") / f"{args.slug}.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(course, indent=2) + "\n")
    print(f"wrote {out_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
